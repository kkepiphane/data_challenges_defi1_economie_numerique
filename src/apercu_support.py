"""
APERCU DU SUPPORT — RENDU PNG DES DIAPOSITIVES
===============================================
Relit le .pptx produit et en dessine une image par diapositive.

POURQUOI
--------
Un support assemble par programme ne se verifie pas en lisant le programme :
deux blocs peuvent se chevaucher, un titre deborder, une bande de chiffres
perdre son alignement, sans qu'aucune ligne de code ne paraisse fausse. Il faut
REGARDER. Sans PowerPoint ni LibreOffice sur le poste, l'aperçu se fabrique en
relisant le fichier et en redessinant ses formes.

CE QUE CET APERCU EST, ET N'EST PAS
------------------------------------
Il reproduit ce que le fichier CONTIENT : positions, tailles, couleurs, corps
de texte, images, tableaux. Il ne reproduit pas le moteur de rendu de
PowerPoint — la cesure et la chasse different de quelques pour cent. Il sert
donc a valider une MISE EN PAGE (equilibre, chevauchements, debordements),
jamais a remplacer une relecture dans PowerPoint avant remise.

Usage :
    python src/apercu_support.py            toutes les diapositives
    python src/apercu_support.py 1 12       seulement celles-la

Sortie : reports/apercu/diapo_NN.png
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / "reports" / "Defi1_Togo_Connectivite_numerique.pptx"
SORTIE = ROOT / "reports" / "apercu"

EMU_PAR_POUCE = 914400
PPP = 110                                   # pixels par pouce de l'aperçu

POLICES = {
    ("Segoe UI", False): r"C:\Windows\Fonts\segoeui.ttf",
    ("Segoe UI", True): r"C:\Windows\Fonts\segoeuib.ttf",
    ("Consolas", False): r"C:\Windows\Fonts\consola.ttf",
    ("Consolas", True): r"C:\Windows\Fonts\consolab.ttf",
}
REPLI = (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf")


def _px(emu) -> float:
    return float(emu) / EMU_PAR_POUCE * PPP


def _police(nom: str, gras: bool, corps_pt: float):
    chemin = POLICES.get((nom, gras)) or REPLI[1 if gras else 0]
    if not Path(chemin).exists():
        chemin = REPLI[1 if gras else 0]
    return ImageFont.truetype(chemin, max(6, int(round(corps_pt * PPP / 72))))


def _couper(dessin, mots, police, largeur, espacement_px):
    """Decoupe un texte en lignes qui tiennent dans `largeur`."""
    def chasse(t):
        return dessin.textlength(t, font=police) + espacement_px * len(t)

    lignes, courante = [], ""
    for mot in mots.split(" "):
        essai = f"{courante} {mot}".strip()
        if courante and chasse(essai) > largeur:
            lignes.append(courante)
            courante = mot
        else:
            courante = essai
    lignes.append(courante)
    return lignes


def _ecrire(dessin, forme):
    cadre = forme.text_frame
    x0, y0 = _px(forme.left), _px(forme.top)
    largeur = _px(forme.width)
    y = y0
    for p in cadre.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            y += 6
            continue
        r = runs[0]
        corps = r.font.size.pt if r.font.size else 12
        police = _police(r.font.name or "Segoe UI", bool(r.font.bold), corps)
        couleur = (17, 17, 17)
        try:
            c = r.font.color.rgb
            couleur = (c[0], c[1], c[2])
        except Exception:
            pass
        # Interlettrage : ecrit dans le XML, donc relu dans le XML.
        spc = r.font._rPr.get("spc")
        espacement = (int(spc) / 100) * PPP / 72 if spc else 0.0
        interligne = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        hauteur = corps * PPP / 72 * interligne

        texte = "".join(run.text for run in runs)
        for ligne in _couper(dessin, texte, police, largeur, espacement):
            x = x0
            if str(p.alignment) == "RIGHT (2)":
                utile = dessin.textlength(ligne, font=police) + espacement * len(ligne)
                x = x0 + largeur - utile
            if espacement:
                for lettre in ligne:                      # lettre a lettre
                    dessin.text((x, y), lettre, font=police, fill=couleur)
                    x += dessin.textlength(lettre, font=police) + espacement
            else:
                dessin.text((x, y), ligne, font=police, fill=couleur)
            y += hauteur


def _tableau(image, dessin, forme):
    tbl = forme.table
    y = _px(forme.top)
    for i, rang in enumerate(tbl.rows):
        x = _px(forme.left)
        h = _px(rang.height)
        for j, cellule in enumerate(rang.cells):
            w = _px(tbl.columns[j].width)
            fond = (255, 255, 255)
            try:
                c = cellule.fill.fore_color.rgb
                fond = (c[0], c[1], c[2])
            except Exception:
                pass
            dessin.rectangle([x, y, x + w, y + h], fill=fond,
                             outline=(230, 229, 224))
            p = cellule.text_frame.paragraphs[0]
            if p.runs:
                r = p.runs[0]
                corps = r.font.size.pt if r.font.size else 9.5
                police = _police(r.font.name or "Segoe UI",
                                 bool(r.font.bold), corps)
                coul = (17, 17, 17)
                try:
                    c = r.font.color.rgb
                    coul = (c[0], c[1], c[2])
                except Exception:
                    pass
                dessin.text((x + 6, y + h / 2), cellule.text, font=police,
                            fill=coul, anchor="lm")
            x += w
        y += h


def _forme(image, dessin, forme):
    if forme.shape_type == MSO_SHAPE_TYPE.PICTURE:
        vignette = Image.open(io.BytesIO(forme.image.blob)).convert("RGB")
        w, h = int(_px(forme.width)), int(_px(forme.height))
        image.paste(vignette.resize((w, h), Image.LANCZOS),
                    (int(_px(forme.left)), int(_px(forme.top))))
        return
    if forme.has_table:
        _tableau(image, dessin, forme)
        return
    if forme.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        boite = [_px(forme.left), _px(forme.top),
                 _px(forme.left) + _px(forme.width),
                 _px(forme.top) + _px(forme.height)]
        remplissage = None
        try:
            if forme.fill.type is not None and forme.fill.type == 1:
                c = forme.fill.fore_color.rgb
                remplissage = (c[0], c[1], c[2])
        except Exception:
            pass
        contour = None
        try:
            if forme.line.fill.type == 1:
                c = forme.line.color.rgb
                contour = (c[0], c[1], c[2])
        except Exception:
            pass
        if remplissage or contour:
            dessin.rectangle(boite, fill=remplissage, outline=contour)
    if forme.has_text_frame and forme.text_frame.text.strip():
        _ecrire(dessin, forme)


def main() -> None:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace")
    if not SUPPORT.exists():
        sys.exit("Support absent : lancer d'abord python src/build_deck.py")

    voulues = {int(a) for a in sys.argv[1:]} or None
    prs = Presentation(str(SUPPORT))
    taille = (int(_px(prs.slide_width)), int(_px(prs.slide_height)))
    SORTIE.mkdir(parents=True, exist_ok=True)

    for n, diapo in enumerate(prs.slides, start=1):
        if voulues and n not in voulues:
            continue
        image = Image.new("RGB", taille, (255, 255, 255))
        dessin = ImageDraw.Draw(image)
        for forme in diapo.shapes:
            try:
                _forme(image, dessin, forme)
            except Exception as e:                 # une forme exotique
                print(f"  diapo {n} : forme ignoree ({e})")
        chemin = SORTIE / f"diapo_{n:02d}.png"
        image.save(chemin)
        print(f"diapo {n:02d} -> {chemin.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
